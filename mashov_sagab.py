from __future__ import annotations

import os
from typing import Any, Dict, List

import matplotlib
import numpy as np

matplotlib.use("TkAgg")
import matplotlib.pyplot as plt
import pandas as pd
from docxtpl import DocxTemplate, InlineImage
from docx import Document
from docx.shared import Inches

from sagab_constants import *

from common_functions import *

def drop_ignored_values(series: pd.Series) -> pd.Series:
    filtered_series = series[series != IGNORED_VALUE]

    return filtered_series


def compute_series_mean_std(series: pd.Series) -> tuple[float, float]:
    numeric_series = clean_numeric_series(series)
    series_real_values = drop_ignored_values(numeric_series)

    if series_real_values.empty:
        return 0.0, 0.0

    mean_val = float(series_real_values.mean())
    std_val = float(series_real_values.std(ddof=1)) if len(series_real_values) > 1 else 0.0

    mean_val = float(rounded_number(mean_val))
    std_val = float(rounded_number(std_val))

    return mean_val, std_val


def compute_numeric_stats_for_df(
    df: pd.DataFrame,
    scope: str,
) -> Dict[str, float]:
    placeholder_to_value: Dict[str, float] = {}

    for question_col in NUMERIC_QUESTION_COLUMNS:
        if question_col not in df.columns:
            continue
        mean_val, std_val = compute_series_mean_std(df[question_col])
        placeholder_map = NUMERIC_PLACEHOLDERS[question_col]

        if scope == NUMERIC_PLACEHOLDER_SCOPE_COMMANDER:
            mean_ph = placeholder_map[NUMERIC_PH_MEAN_COMMANDER]
            std_ph = placeholder_map[NUMERIC_PH_STD_COMMANDER]
            placeholder_to_value[mean_ph] = mean_val
            placeholder_to_value[std_ph] = std_val

        elif scope == NUMERIC_PLACEHOLDER_SCOPE_COHORT:
            cohort_mean_ph = placeholder_map[NUMERIC_PH_MEAN_COHORT]
            placeholder_to_value[cohort_mean_ph] = mean_val

        else:
            raise ValueError(f"Unknown scope: {scope}")

    return placeholder_to_value


def compute_commander_numeric_stats(df_commander: pd.DataFrame) -> Dict[str, float]:
    return compute_numeric_stats_for_df(
        df=df_commander,
        scope=NUMERIC_PLACEHOLDER_SCOPE_COMMANDER,
    )


def compute_cohort_numeric_stats(df_all: pd.DataFrame) -> Dict[str, float]:
    return compute_numeric_stats_for_df(
        df=df_all,
        scope=NUMERIC_PLACEHOLDER_SCOPE_COHORT,
    )


def create_all_section_charts(
    placeholder_to_value: Dict[str, Any],
    output_dir: str,
) -> Dict[str, str]:
    os.makedirs(output_dir, exist_ok=True)
    section_prefix_to_path: Dict[str, str] = {}

    for section_label, section_prefix in NUMERIC_SECTION_TO_PREFIX.items():
        img_path = create_section_bar_chart(
            section_label=section_label,
            section_prefix=section_prefix,
            placeholder_to_value=placeholder_to_value,
            output_dir=output_dir,
        )
        section_prefix_to_path[section_prefix] = img_path

    return section_prefix_to_path



def _compute_x_positions(num_questions: int) -> List[float]:
    if num_questions <= 0:
        return []
    if num_questions >= MAX_QUESTIONS_PER_SECTION:
        return list(range(MAX_QUESTIONS_PER_SECTION))
    return list(np.linspace(0, MAX_QUESTIONS_PER_SECTION - 1, num_questions))


def compute_centered_positions(num_questions: int) -> List[float]:
    if num_questions <= 0:
        return []

    left = 0.0
    right = float(MAX_QUESTIONS_PER_SECTION - 1)
    center = (left + right) / 2.0  # here: 1.0

    if num_questions >= MAX_QUESTIONS_PER_SECTION:
        return [left + i for i in range(MAX_QUESTIONS_PER_SECTION)]

    if num_questions == 1:
        return [center]

    if num_questions == 2:
        half_spacing = 0.5
        return [center - half_spacing, center + half_spacing]

    step = (right - left) / float(num_questions - 1)
    return [left + i * step for i in range(num_questions)]
def create_section_bar_chart(
    section_label: str,
    section_prefix: str,
    placeholder_to_value: Dict[str, Any],
    output_dir: str,
    figure_width: float = 10.0,
    figure_height: float = 7.0,
) -> str:
    commander_means, commander_stds, cohort_means, x_labels = _collect_section_stats(
        section_label,
        placeholder_to_value,
    )

    fig = plt.figure(figsize=(figure_width, figure_height))
    ax = fig.add_subplot(111)

    _plot_section_bars(
        ax=ax,
        commander_means=commander_means,
        commander_stds=commander_stds,
        cohort_means=cohort_means,
    )

    _style_section_axes(
        ax=ax,
        section_label=section_label,
        x_labels=x_labels,
    )

    fig.subplots_adjust(
        left=0.08,
        right=0.98,
        top=0.9,
        bottom=0.25,
    )

    filename = f"numeric_{section_prefix}.png"
    file_path = os.path.join(output_dir, filename)
    fig.savefig(file_path, dpi=200)
    plt.close(fig)

    return file_path


def _style_section_axes(
    ax: plt.Axes,
    section_label: str,
    x_labels: List[str],
) -> None:
    num_questions = len(x_labels)
    x_positions = compute_centered_positions(num_questions)

    ax.set_title(rtl_embed_graphic(section_label))
    ax.set_ylabel(rtl_embed_graphic("ממוצע"))
    ax.set_ylim(1, 6)
    ax.set_yticks(range(1, 7))
    short_labels = [SHORT_LABELS.get(label, label) for label in x_labels]
    display_labels = [rtl_embed_graphic(lbl) for lbl in short_labels]

    ax.set_xticks(x_positions)
    ax.set_xticklabels(display_labels, ha="center")

    ax.set_xlim(-0.5, float(MAX_QUESTIONS_PER_SECTION - 0.5))

    ax.legend(
        loc="upper center",
        bbox_to_anchor=(0.5, -0.15),
        ncol=2,
        fontsize=14,
        frameon=True,
        framealpha=1.0,
    )

    ax.grid(axis="y", linestyle="--", alpha=0.3)






def _get_section_questions(section_label: str) -> List[str]:
    if section_label not in NUMERIC_SECTION_TO_QUESTIONS:
        raise ValueError(f"Unknown section label: {section_label}")
    return NUMERIC_SECTION_TO_QUESTIONS[section_label]


def _collect_section_stats(
    section_label: str,
    placeholder_to_value: Dict[str, Any],
) -> tuple[List[float], List[float], List[float], List[str]]:
    questions = _get_section_questions(section_label)

    commander_means: List[float] = []
    commander_stds: List[float] = []
    cohort_means: List[float] = []
    x_labels: List[str] = []

    for question_col in questions:
        mean_cmd, std_cmd, mean_cohort = _get_question_stats_from_placeholders(
            question_col,
            placeholder_to_value,
        )
        commander_means.append(mean_cmd)
        commander_stds.append(std_cmd)
        cohort_means.append(mean_cohort)
        x_labels.append(question_col)

    return commander_means, commander_stds, cohort_means, x_labels


def _get_question_stats_from_placeholders(
    question_col: str,
    placeholder_to_value: Dict[str, Any],
) -> tuple[float, float, float]:
    placeholder_map = NUMERIC_PLACEHOLDERS[question_col]

    mean_cmd_name = placeholder_map[NUMERIC_PH_MEAN_COMMANDER]
    std_cmd_name = placeholder_map[NUMERIC_PH_STD_COMMANDER]
    mean_cohort_name = placeholder_map[NUMERIC_PH_MEAN_COHORT]

    mean_cmd_val = float(placeholder_to_value.get(mean_cmd_name, 0.0))
    std_cmd_val = float(placeholder_to_value.get(std_cmd_name, 0.0))
    mean_cohort_val = float(placeholder_to_value.get(mean_cohort_name, 0.0))

    return mean_cmd_val, std_cmd_val, mean_cohort_val


def _plot_section_bars(
    ax: plt.Axes,
    commander_means: List[float],
    commander_stds: List[float],
    cohort_means: List[float],
    bar_width: float = 0.35,
) -> None:
    num_questions = len(commander_means)
    x_positions = compute_centered_positions(num_questions)

    ax.bar(
        [x - bar_width / 2 for x in x_positions],
        commander_means,
        width=bar_width,
        yerr=commander_stds,
        capsize=5,
        label=rtl_embed_graphic("ממוצע אישי"),
    )

    ax.bar(
        [x + bar_width / 2 for x in x_positions],
        cohort_means,
        width=bar_width,
        label=rtl_embed_graphic("ממוצע סגל"),
        alpha=0.8,
    )





def build_graphs_context(
    tpl: DocxTemplate,
    section_prefix_to_path: Dict[str, str],
) -> Dict[str, Any]:
    context: Dict[str, Any] = {}

    for section_prefix, img_path in section_prefix_to_path.items():
        placeholder_name = f"chart_{section_prefix}"
        context[placeholder_name] = InlineImage(
            tpl,
            img_path,
            width=Inches(5.5),
        )

    return context


def build_template_context(
    tpl: DocxTemplate,
    df_commander: pd.DataFrame,
    commander: str,
    charts_dir: str,
    placeholder_to_value: Dict[str, Any],
) -> Dict[str, Any]:
    charts_paths = create_all_section_charts(
        placeholder_to_value=placeholder_to_value,
        output_dir=charts_dir,
    )

    graphs_context = build_graphs_context(tpl, charts_paths)
    bullets_context = merge_bullet_lists(
        df_commander=df_commander,
        commander=commander,
        bullet_list_context=BULLET_LIST_CONTEXT,
    )



    context: Dict[str, Any] = {}
    context.update(graphs_context)
    context.update(bullets_context)

    return context


def validate_calculations(placeholder_to_value: Dict[str, Any]) -> bool:
    if None in placeholder_to_value.values() or "" in placeholder_to_value.values():
        print("Some placeholders have empty values.")
        return False
    return True


def generate_commander_docx(
    df_all: pd.DataFrame,
    df_commander: pd.DataFrame,
    commander: str,
    placeholder_to_value: Dict[str, Any],
) -> None:
    os.makedirs(OUTPUT_PATH, exist_ok=True)
    charts_dir = os.path.join(OUTPUT_PATH, "graphs")

    doc_path = os.path.join(OUTPUT_PATH, f"{commander}.docx")

    doc = Document(TEMPLATE_PATH)
    replace_placeholders(doc, placeholder_to_value)
    doc.save(doc_path)

    tpl = DocxTemplate(doc_path)
    context = build_template_context(
        tpl=tpl,
        df_commander=df_commander,
        commander=commander,
        charts_dir=charts_dir,
        placeholder_to_value=placeholder_to_value,
    )
    tpl.render(context)
    tpl.save(doc_path)


def main(file_path: str = INPUT_PATH) -> None:
    try:
        df = excel_to_dataframe(file_path)
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return

    if not validate_excel(df, COLUMNS):
        print("Excel validation failed.")
        return

    df = clean_dataframe(df, COMMANDER_COLUMN)

    cohort_numeric_stats = compute_cohort_numeric_stats(df)

    for commander in df[COMMANDER_COLUMN].unique():
        df_commander = df[df[COMMANDER_COLUMN] == commander]

        commander_numeric_stats = compute_commander_numeric_stats(df_commander)

        placeholder_to_value: Dict[str, Any] = {}
        placeholder_to_value.update(commander_numeric_stats)
        placeholder_to_value.update(cohort_numeric_stats)

        if not validate_calculations(placeholder_to_value):
            print(f"Calculations validation failed for commander {commander}.")
            continue

        generate_commander_docx(
            df_all=df,
            df_commander=df_commander,
            commander=commander,
            placeholder_to_value=placeholder_to_value,
        )


if __name__ == "__main__":
    main()
